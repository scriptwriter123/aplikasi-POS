import os
from datetime import date, datetime, time, timedelta, timezone

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
import pytest
from PySide6.QtCore import QEvent
from PySide6.QtWidgets import QApplication, QMessageBox
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.core.database import Base
from app.hardware.report_exporter import ReportExporter
from app.models import (
    CashierDailyReport,
    Product,
    Role,
    Sale,
    SaleItem,
    Shift,
    Store,
    User,
)
from app.services.cashier_daily_report_service import CashierDailyReportService
from app.services.report_periods import calculate_period_range
from app.services.report_service import ReportService
from app.ui.views.report_view import (
    CashierDailyReportDetailsDialog,
    ReportView,
)


@pytest.fixture
def report_session():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    try:
        with Session(engine) as session:
            yield session
    finally:
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_period_calculation_supports_hour_day_month_and_year():
    sample = datetime(2026, 10, 9, 14, 37)

    hour_start, hour_end, hour_label = calculate_period_range("hour", sample)
    assert hour_start.hour == 14
    assert hour_start.minute == 0
    assert (hour_end - hour_start).total_seconds() == 3600
    assert hour_label.endswith("14:00")

    day_start, day_end, day_label = calculate_period_range("day", sample.date())
    assert day_start.date() == sample.date()
    assert day_end.date().day == 10
    assert day_label == "09-10-2026"

    month_start, month_end, month_label = calculate_period_range(
        "month",
        sample.date(),
    )
    assert month_start.day == 1
    assert month_end.month == 11
    assert month_label == "October 2026"

    year_start, year_end, year_label = calculate_period_range("year", sample)
    assert year_start.month == 1 and year_start.day == 1
    assert year_end.year == 2027
    assert year_label == "2026"


def test_report_service_uses_saved_item_cost_and_marks_missing_cost(
    report_session: Session,
):
    role = Role(name="CASHIER")
    store = Store(code="MAIN", name="Test Store")
    report_session.add_all([role, store])
    report_session.flush()
    cashier = User(
        username="cashier",
        full_name="Test Cashier",
        password_hash="unused",
        role=role,
    )
    report_session.add(cashier)
    report_session.flush()
    shift = Shift(
        store_id=store.id,
        cashier_id=cashier.id,
        opened_at=datetime(2026, 10, 9, 0, tzinfo=timezone.utc),
    )
    report_session.add(shift)
    report_session.flush()
    product = Product(sku="SKU-1", name="Product", purchase_price=200)
    report_session.add(product)
    report_session.flush()
    sale = Sale(
        invoice_no="INV-1",
        store_id=store.id,
        shift_id=shift.id,
        cashier_id=cashier.id,
        subtotal=1_100,
        grand_total=1_100,
        created_at=datetime(2026, 10, 9, 8, tzinfo=timezone.utc),
    )
    report_session.add(sale)
    report_session.flush()
    report_session.add_all(
        [
            SaleItem(
                sale_id=sale.id,
                product_id=product.id,
                qty=2,
                unit_price=500,
                unit_cost=200,
                line_total=1_000,
            ),
            SaleItem(
                sale_id=sale.id,
                product_id=product.id,
                qty=1,
                unit_price=100,
                unit_cost=None,
                line_total=100,
            ),
        ]
    )
    report_session.flush()

    service = ReportService(report_session)
    start = datetime(2026, 10, 9, 0, tzinfo=timezone.utc)
    end = datetime(2026, 10, 10, 0, tzinfo=timezone.utc)
    sales = service.sales(start, end)
    profits = service.profits(start, end)

    assert len(sales) == 1
    assert sales[0]["revenue"] == 1_100
    assert sales[0]["items"] == [
        {
            "product_name": "Product",
            "product_sku": "SKU-1",
            "quantity": 2,
            "unit_price": 500,
            "line_total": 1_000,
        },
        {
            "product_name": "Product",
            "product_sku": "SKU-1",
            "quantity": 1,
            "unit_price": 100,
            "line_total": 100,
        },
    ]
    assert len(profits) == 1
    assert profits[0]["known_cost"] == 400
    assert profits[0]["known_profit"] == 600
    assert profits[0]["missing_cost_count"] == 1


def test_low_stock_report_lists_active_products_below_five_only(
    report_session: Session,
):
    products = [
        Product(sku="LOW-0", name="Out of stock", stock_qty=0),
        Product(sku="LOW-4", name="Low stock", stock_qty=4),
        Product(sku="EDGE-5", name="At threshold", stock_qty=5),
        Product(
            sku="INACTIVE",
            name="Inactive low stock",
            stock_qty=2,
            is_active=False,
        ),
        Product(
            sku="DELETED",
            name="Deleted low stock",
            stock_qty=1,
            deleted_at=datetime.now(timezone.utc),
        ),
    ]
    report_session.add_all(products)
    report_session.flush()

    low_stock = ReportService(report_session).low_stock_products()

    assert [product.sku for product in low_stock] == ["LOW-0", "LOW-4"]

    app = QApplication.instance() or QApplication([])
    view = ReportView(report_session, "sales")
    view._populate_low_stock_products()
    label_text = view.low_stock_label.text()
    assert "Stok habis (0):" in label_text
    assert "Out of stock (LOW-0)" in label_text
    assert "Stok rendah (< 5):" in label_text
    assert "Low stock (LOW-4): stok 4" in label_text
    view.close()


def test_report_can_be_saved_to_pdf_and_docx(tmp_path):
    app = QApplication.instance() or QApplication([])
    exporter = ReportExporter()
    report = {
        "title": "Laporan Penjualan",
        "period": "09-10-2026",
        "generated_at": "09-10-2026 12:00:00",
        "summary": "Total penjualan: Rp1.000",
        "warning": "",
        "headers": ["Invoice", "Total"],
        "rows": [["INV-1", "Rp1.000"]],
    }

    pdf_path = exporter.save_pdf(report, tmp_path / "laporan.pdf")
    docx_path = exporter.save_docx(report, tmp_path / "laporan.docx")
    app.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    app.processEvents()

    assert pdf_path.read_bytes().startswith(b"%PDF-")
    assert docx_path.read_bytes().startswith(b"PK")
    word_document = Document(docx_path)
    assert word_document.paragraphs[0].text == "Laporan Penjualan"
    assert word_document.tables[0].rows[1].cells[0].text == "INV-1"


def test_report_printer_filter_excludes_thermal_devices():
    assert ReportExporter.is_thermal_printer("XPrinter XP-80")
    assert ReportExporter.is_thermal_printer("Epson TM-T88")
    assert not ReportExporter.is_thermal_printer("HP LaserJet Office")


def test_report_service_rejects_reversed_period(report_session: Session):
    service = ReportService(report_session)
    with pytest.raises(ValueError, match="sebelum"):
        service.sales(
            datetime(2026, 10, 10, tzinfo=timezone.utc),
            datetime(2026, 10, 9, tzinfo=timezone.utc),
        )


def test_cashier_daily_report_saves_totals_and_updates_existing_report(
    report_session: Session,
):
    role = Role(name="CASHIER")
    store = Store(code="MAIN", name="Test Store")
    report_session.add_all([role, store])
    report_session.flush()
    cashier = User(
        username="cashier",
        full_name="Test Cashier",
        password_hash="unused",
        role=role,
    )
    report_session.add(cashier)
    report_session.flush()
    shift = Shift(
        store_id=store.id,
        cashier_id=cashier.id,
        opened_at=datetime(2026, 10, 9, tzinfo=timezone.utc),
    )
    report_session.add(shift)
    report_session.flush()
    report_date = date(2026, 10, 9)
    sale_time = datetime.combine(report_date, time(12)).astimezone(
        timezone.utc
    ).replace(tzinfo=None)
    report_session.add(
        Sale(
            invoice_no="INV-DAILY-1",
            store_id=store.id,
            shift_id=shift.id,
            cashier_id=cashier.id,
            subtotal=25_000,
            grand_total=25_000,
            status="PAID",
            created_at=sale_time,
        )
    )
    report_session.flush()
    product = Product(
        sku="SALE-ITEM",
        name="Kopi",
        selling_price=25_000,
        stock_qty=3,
    )
    report_session.add(product)
    report_session.flush()
    sale = report_session.query(Sale).filter_by(invoice_no="INV-DAILY-1").one()
    report_session.add(
        SaleItem(
            sale_id=sale.id,
            product_id=product.id,
            qty=1,
            unit_price=25_000,
            line_total=25_000,
        )
    )
    report_session.flush()

    service = CashierDailyReportService(report_session)
    saved = service.save(cashier, report_date, "  Stok minuman menipis.  ")
    report_id = saved.id
    assert saved.sales_count == 1
    assert saved.sales_total == 25_000
    assert saved.notes == "Stok minuman menipis."

    updated = service.save(cashier, report_date, "Sudah diperbarui.")
    assert updated.id == report_id
    assert updated.notes == "Sudah diperbarui."
    assert report_session.query(CashierDailyReport).count() == 1

    start = datetime.combine(report_date, time.min).astimezone(timezone.utc)
    end = datetime.combine(
        report_date + timedelta(days=1),
        time.min,
    ).astimezone(timezone.utc)
    reports = service.reports(start, end)
    assert len(reports) == 1
    assert reports[0].cashier.full_name == "Test Cashier"

    app = QApplication.instance() or QApplication([])
    view = ReportView(report_session, "sales")
    assert view.period_stack.width() == view.period_mode.width()
    assert not view.back_button.icon().isNull()
    assert view.report_table.isHidden()
    assert "Harian" in view.heading_label.text()
    view.period_mode.setCurrentIndex(1)
    assert "Bulanan" in view.heading_label.text()
    view.period_mode.setCurrentIndex(2)
    assert "Tahunan" in view.heading_label.text()
    view.period_mode.setCurrentIndex(0)
    view.daily_date.setDate(view.daily_date.date().addDays(-1))
    assert view.daily_date.date().toString("dd-MM-yyyy") in view.heading_label.text()
    view._populate_sales_report(
        ReportService(report_session).sales(start, end),
        "09-10-2026",
    )
    assert view.report_table.item(0, 3).text() == "Kopi"
    assert view.report_table.item(0, 6).text() == "Rp 25.000,00"
    assert view.report_table.item(0, 7).text() == "Rp 25.000,00"
    view._populate_cashier_daily_reports(start, end)
    assert view.cashier_reports_table.rowCount() == 1
    assert view.cashier_reports_table.item(0, 1).text() == "Test Cashier"
    assert view.cashier_reports_table.item(0, 3).text() == "Rp 25.000,00"
    assert view.cashier_reports_table.item(0, 4).text() == "Sudah diperbarui."
    details_button = view.cashier_reports_table.cellWidget(0, 5)
    assert details_button is not None
    assert details_button.text() == "Rincian Transaksi"
    assert "#1d4ed8" in details_button.styleSheet()

    cashier_sales = ReportService(report_session).sales(
        start,
        end,
        cashier_id=cashier.id,
    )

    class FakeExporter:
        def __init__(self):
            self.report_data = None

        def print_report(self, report_data):
            self.report_data = report_data
            return "Printer Dokumen"

    exporter = FakeExporter()
    details_dialog = CashierDailyReportDetailsDialog(
        updated,
        cashier_sales,
        "09-10-2026",
        exporter,
        view,
    )
    assert details_dialog.details_table.rowCount() == 1
    assert details_dialog.details_table.item(0, 0).text() == "INV-DAILY-1"
    assert details_dialog.details_table.item(0, 2).text() == "Kopi"
    assert details_dialog.notes_label.text() == (
        "Catatan kasir: Sudah diperbarui."
    )
    monkeypatch = pytest.MonkeyPatch()
    monkeypatch.setattr(
        QMessageBox,
        "information",
        lambda *_args: None,
    )
    details_dialog.print_button.click()
    assert exporter.report_data is not None
    assert exporter.report_data["warning"] == (
        "Catatan kasir: Sudah diperbarui."
    )
    assert exporter.report_data["rows"][0][0] == "INV-DAILY-1"
    details_dialog.back_button.click()
    assert details_dialog.result() == details_dialog.DialogCode.Rejected
    monkeypatch.undo()
    low_stock_product = Product(
        sku="LOW-REPORT",
        name="Produk stok tipis",
        stock_qty=3,
    )
    report_session.add(low_stock_product)
    report_session.flush()
    view._populate_low_stock_products()
    assert "Produk stok tipis (LOW-REPORT): stok 3" in view.low_stock_label.text()
    view.close()
    profit_view = ReportView(report_session, "profit")
    assert not profit_view.report_table.isHidden()
    profit_view.close()
    app.processEvents()
